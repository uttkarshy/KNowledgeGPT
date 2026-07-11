"""PPTX extraction via python-pptx. Each slide's title becomes the section
title for all content on that slide; slide number is treated as page_number
so citations can say "Slide 4"."""

from __future__ import annotations

from pptx import Presentation

from app.core.config import Settings
from app.services.extraction.base import TextExtractor
from app.services.extraction.schemas import (
    ExtractedBlock,
    ExtractedDocument,
    ExtractedPage,
    ExtractionError,
    SectionKind,
)


class PPTXExtractor(TextExtractor):
    def extract(self, file_path: str, *, settings: Settings) -> ExtractedDocument:
        try:
            prs = Presentation(file_path)
        except Exception as e:
            raise ExtractionError(f"Could not open PPTX: {e}") from e

        pages: list[ExtractedPage] = []

        for slide_index, slide in enumerate(prs.slides):
            slide_number = slide_index + 1
            blocks: list[ExtractedBlock] = []
            slide_title: str | None = None

            title_shape = slide.shapes.title
            if title_shape is not None and title_shape.has_text_frame:
                title_text = title_shape.text_frame.text.strip()
                if title_text:
                    slide_title = title_text
                    blocks.append(
                        ExtractedBlock(
                            kind=SectionKind.HEADING, text=title_text, heading_level=1,
                            page_number=slide_number, section_title=slide_title,
                        )
                    )

            for shape in slide.shapes:
                if shape == title_shape or not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        blocks.append(
                            ExtractedBlock(
                                kind=SectionKind.PARAGRAPH, text=text,
                                page_number=slide_number, section_title=slide_title,
                            )
                        )

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                blocks.append(
                    ExtractedBlock(
                        kind=SectionKind.PARAGRAPH,
                        text=f"[Speaker notes] {slide.notes_slide.notes_text_frame.text.strip()}",
                        page_number=slide_number, section_title=slide_title,
                    )
                )

            pages.append(ExtractedPage(page_number=slide_number, blocks=blocks, was_ocr=False))

        return ExtractedDocument(pages=pages, detected_language=None, page_count=len(pages))
