import zipfile

import openpyxl
import pytest
from docx import Document

from app.core.config import Settings
from app.services.extraction.csv_extractor import CSVExtractor
from app.services.extraction.docx_extractor import DOCXExtractor
from app.services.extraction.schemas import ExtractionError
from app.services.extraction.text_markdown_extractor import TextMarkdownExtractor
from app.services.extraction.xlsx_extractor import XLSXExtractor
from app.services.upload.validation import FileValidationError, validate_uploaded_file


def text(result):
    return '\n'.join(block.text for page in result.pages for block in page.blocks)


def validate(path):
    return validate_uploaded_file(file_path=str(path), filename=path.name, size_bytes=path.stat().st_size, settings=Settings())


def test_docx_txt_csv_real_extraction(tmp_path):
    doc = Document()
    doc.add_paragraph('Full document evidence')
    path = tmp_path / 'example.docx'
    doc.save(path)
    validate(path)
    assert 'Full document evidence' in text(DOCXExtractor().extract(str(path), settings=Settings()))
    for extension, extractor, payload in [('txt', TextMarkdownExtractor(is_markdown=False), 'Full text evidence'),
                                           ('csv', CSVExtractor(), 'date,amount\n31 Jul 2026,309.00')]:
        path = tmp_path / f'example.{extension}'
        path.write_text(payload)
        validate(path)
        assert payload.splitlines()[-1].replace(',', ' | ') in text(extractor.extract(str(path), settings=Settings()))


@pytest.mark.parametrize('extension', ['txt', 'csv', 'md'])
@pytest.mark.parametrize('payload', [b'MZ' + b'\x00' * 100, b'%PDF-1.7\n', b'hello\x00binary'])
def test_disguised_binary_rejected(tmp_path, extension, payload):
    path = tmp_path / f'fake.{extension}'
    path.write_bytes(payload)
    with pytest.raises(FileValidationError):
        validate(path)


def test_archive_expansion_guard(tmp_path):
    path = tmp_path / 'bomb.docx'
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', b'0' * (100 * 1024 * 1024 + 1))
    with pytest.raises(FileValidationError, match='safe extraction limits'):
        validate(path)


def test_large_real_spreadsheet_retains_last_row_and_rejects_overflow(tmp_path):
    path = tmp_path / 'large.xlsx'
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(['Row', 'Amount'])
    for i in range(4999):
        sheet.append([f'transaction-{i}', i])
    workbook.save(path)
    validate(path)
    result = XLSXExtractor().extract(str(path), settings=Settings())
    assert 'transaction-4998' in text(result)
    assert 'transaction-0' in text(result)
    sheet.append(['overflow', 999])
    workbook.save(path)
    with pytest.raises(ExtractionError, match='split the file'):
        XLSXExtractor().extract(str(path), settings=Settings())
    workbook.close()


def test_csv_overflow_is_explicit(tmp_path):
    path = tmp_path / 'large.csv'
    path.write_text('date,amount\n' + '31 Jul 2026,1\n' * 10000)
    with pytest.raises(ExtractionError, match='split the file'):
        CSVExtractor().extract(str(path), settings=Settings())

